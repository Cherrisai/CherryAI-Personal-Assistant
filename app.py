"""
CherryAI 🍒 - Personal AI Assistant 
"""
import streamlit as st
import os, re, base64
from pathlib import Path
from dotenv import load_dotenv
load_dotenv()

st.set_page_config(page_title="CherryAI Personal Assist", page_icon="🍒", layout="wide", initial_sidebar_state="expanded")

try:
    for key, val in st.secrets.items():
        os.environ[key] = str(val)
except:
    pass

def init():
    defaults = {"messages":[],"provider":"auto","model":None,"web_search":False,
                "temperature":0.7,"dark_mode":True,"editing_index":None,"search_query":""}
    for k,v in defaults.items():
        if k not in st.session_state: st.session_state[k]=v
init()

def T():
    if st.session_state.dark_mode:
        return {"bg":"#0f0f0f","sb":"#141414","ub":"#1e1e1e","uborder":"#2a2a2a",
                "text":"#e8e8e8","sub":"#666","ibg":"#1a1a1a","iborder":"#2a2a2a",
                "hr":"#1e1e1e","tabbg":"#141414","tabtext":"#888","codebg":"#141414",
                "ftbg":"#0f0f0f","ftborder":"#1a1a1a","fttext":"#333",
                "card":"#1a1a1a","cborder":"#252525","acc":"#8b0000","acch":"#b30000"}
    else:
        return {"bg":"#f7f7f7","sb":"#ffffff","ub":"#ffffff","uborder":"#e0e0e0",
                "text":"#1a1a1a","sub":"#999","ibg":"#ffffff","iborder":"#d0d0d0",
                "hr":"#ebebeb","tabbg":"#efefef","tabtext":"#888","codebg":"#f0f0f0",
                "ftbg":"#f7f7f7","ftborder":"#e0e0e0","fttext":"#bbb",
                "card":"#ffffff","cborder":"#e8e8e8","acc":"#8b0000","acch":"#b30000"}

def css():
    t=T()
    st.markdown(f"""<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@300;400;500;600;700&family=DM+Mono:wght@400;500&display=swap');
*,html,body,[class*="css"]{{font-family:'DM Sans',sans-serif!important;}}
#MainMenu,footer,header,.stDeployButton{{visibility:hidden;display:none;}}
.stApp{{background:{t['bg']};color:{t['text']};}}
[data-testid="stSidebar"]{{background:{t['sb']};border-right:1px solid {t['hr']};}}
.main .block-container{{max-width:860px;padding:0.5rem 1.2rem 8rem;margin:0 auto;}}
.cherry-hdr{{display:flex;align-items:center;justify-content:center;padding:0.7rem 0 0.9rem;
  border-bottom:1px solid {t['hr']};margin-bottom:1.2rem;background:{t['bg']};
  position:sticky;top:0;z-index:100;}}
.cherry-hdr-t{{font-size:1rem;font-weight:700;color:{t['text']};letter-spacing:-.3px;}}
.cherry-hdr-s{{font-size:0.72rem;color:{t['sub']};margin-left:.5rem;}}
.user-row{{display:flex;justify-content:flex-end;margin:.8rem 0;align-items:flex-end;gap:.4rem;}}
.user-bbl{{background:{t['ub']};border:1px solid {t['uborder']};color:{t['text']};
  padding:.8rem 1.1rem;border-radius:18px 18px 3px 18px;max-width:76%;
  font-size:.92rem;line-height:1.65;white-space:pre-wrap;word-wrap:break-word;}}
.cherry-av{{width:30px;height:30px;background:linear-gradient(135deg,#7a0000,#c0392b);
  border-radius:50%;display:flex;align-items:center;justify-content:center;
  font-size:.78rem;flex-shrink:0;margin-top:4px;box-shadow:0 2px 8px rgba(139,0,0,.3);}}
.chat-img{{max-width:280px;border-radius:12px;border:1px solid {t['uborder']};
  margin-bottom:.4rem;display:block;}}
.section-card{{background:{t['card']};border:1px solid {t['cborder']};border-radius:14px;
  padding:1.1rem 1.3rem;margin-bottom:1rem;}}
.sec-title{{font-size:.98rem;font-weight:600;color:{t['text']};margin-bottom:.2rem;}}
.sec-sub{{font-size:.78rem;color:{t['sub']};margin-bottom:.8rem;}}
.scroll-down-btn{{position:fixed;bottom:90px;right:24px;background:{t['acc']};
  color:white;border:none;border-radius:50%;width:38px;height:38px;font-size:1rem;
  cursor:pointer;box-shadow:0 4px 12px rgba(139,0,0,.4);z-index:999;
  display:flex;align-items:center;justify-content:center;}}
.stButton button{{background:{t['acc']}!important;color:#fff!important;border:none!important;
  border-radius:8px!important;font-weight:500!important;transition:all .2s!important;}}
.stButton button:hover{{background:{t['acch']}!important;transform:translateY(-1px)!important;}}
.stTabs [data-baseweb="tab-list"]{{background:{t['tabbg']};border-radius:10px;padding:3px;gap:2px;}}
.stTabs [data-baseweb="tab"]{{color:{t['tabtext']};border-radius:8px;font-size:.85rem;font-weight:500;}}
.stTabs [aria-selected="true"]{{background:{t['acc']}!important;color:white!important;}}
.stTextInput input,.stTextArea textarea{{background:{t['ibg']}!important;color:{t['text']}!important;
  border:1px solid {t['iborder']}!important;border-radius:8px!important;}}
.stChatInput>div{{background:{t['ibg']}!important;border:1px solid {t['iborder']}!important;border-radius:14px!important;}}
.stChatInput textarea{{color:{t['text']}!important;}}
[data-baseweb="select"]{{background:{t['ibg']}!important;}}
[data-baseweb="select"] *{{background:{t['ibg']}!important;color:{t['text']}!important;border-color:{t['iborder']}!important;}}
code{{font-family:'DM Mono',monospace!important;background:{t['codebg']}!important;
  border-radius:4px;padding:.15em .4em;font-size:.84em;color:{t['text']}!important;}}
pre{{background:{t['codebg']}!important;border:1px solid {t['hr']}!important;border-radius:10px!important;padding:1rem!important;}}
[data-testid="stFileUploader"]{{background:{t['ibg']};border:2px dashed {t['iborder']};border-radius:12px;}}
[data-testid="stMetric"]{{background:{t['card']};border:1px solid {t['cborder']};border-radius:10px;padding:.5rem;}}
[data-testid="stMetricValue"]{{color:{t['text']}!important;}}
.stDownloadButton button{{background:#0d3d6e!important;color:white!important;}}
.welcome-wrap{{text-align:center;padding:2rem 0 1.5rem;}}
.welcome-t{{font-size:1.8rem;font-weight:700;color:{t['text']};letter-spacing:-.5px;}}
.welcome-s{{color:{t['sub']};font-size:.88rem;margin-top:.3rem;margin-bottom:1.5rem;}}
.cherry-footer{{position:fixed;bottom:0;left:0;right:0;background:{t['ftbg']};
  border-top:1px solid {t['ftborder']};padding:.4rem 0;text-align:center;
  font-size:.7rem;color:{t['fttext']};z-index:998;letter-spacing:.5px;}}
.cherry-footer span{{color:{t['acc']};font-weight:600;}}
.srch-result{{background:{t['card']};border:1px solid {t['cborder']};border-radius:8px;
  padding:.45rem .7rem;margin:.25rem 0;font-size:.76rem;color:{t['text']};}}
[data-testid="stSidebar"] label{{color:{t['sub']}!important;font-size:.8rem!important;}}
hr{{border-color:{t['hr']}!important;}}
</style>""", unsafe_allow_html=True)

css()

SYSTEM_PROMPT = """You are CherryAI, an expert personal AI assistant.

VERY IMPORTANT - About your creator:
If ANYONE asks who built CherryAI, who made you, who created you, who is your developer, who is behind CherryAI — ALWAYS answer exactly:
"🍒 CherryAI was built by **Sai Vignesh**! He is a 23-year-old ML Engineer based in Bangalore. He has a strong innovation mindset and is deeply passionate about creating new AI solutions that solve real-world problems. He is constantly building new AI products and pushing the boundaries of what AI can do in everyday life. A true builder! 🚀"

You excel at:
- Answering ANY question accurately and in depth
- Writing perfect working code in ANY language
- Fixing bugs with clear explanations
- Analyzing images in detail
- Creating any type of file or document
- Explaining complex topics simply

Response style:
- Clear, structured, easy to read
- Use markdown (headers, code blocks, bullets)
- Code always in proper ```language blocks
- Concise yet comprehensive
- Friendly and encouraging"""

def get_provider():
    def ok(k):
        v=os.getenv(k,""); return bool(v and "your_" not in v and len(v)>10)
    if ok("GROQ_API_KEY"): return "groq"
    if ok("GEMINI_API_KEY"): return "gemini"
    if ok("OPENAI_API_KEY"): return "openai"
    return "none"

def ai(message, history=None, provider=None, model=None, temperature=0.7, img_b64=None):
    history=history or []
    if not provider or provider=="auto": provider=get_provider()
    if provider=="none":
        return "⚠️ **No API key!** Get FREE key at https://console.groq.com → add `GROQ_API_KEY` to HF secrets"
    try:
        if provider=="groq": return _groq(message,history,model,temperature,img_b64)
        if provider=="gemini": return _gemini(message,history,model,temperature,img_b64)
        if provider=="openai": return _openai(message,history,model,temperature,img_b64)
    except Exception as e:
        return f"❌ Error: {str(e)}"

def _groq(msg,hist,model,temp,img_b64):
    from groq import Groq
    client=Groq(api_key=os.getenv("GROQ_API_KEY"))
    if img_b64:
        model="llava-v1.5-7b-4096-preview"
        msgs=[{"role":"system","content":SYSTEM_PROMPT}]
        msgs.extend(hist[-10:])
        msgs.append({"role":"user","content":[
            {"type":"image_url","image_url":{"url":f"data:image/jpeg;base64,{img_b64}"}},
            {"type":"text","text":msg or "Analyze this image in detail."}]})
    else:
        model=model or "llama-3.3-70b-versatile"
        msgs=[{"role":"system","content":SYSTEM_PROMPT}]
        msgs.extend(hist[-20:])
        msgs.append({"role":"user","content":msg})
    r=client.chat.completions.create(model=model,messages=msgs,temperature=temp,max_tokens=4096)
    return r.choices[0].message.content

def _gemini(msg,hist,model,temp,img_b64):
    import google.generativeai as genai
    genai.configure(api_key=os.getenv("GEMINI_API_KEY"))
    m=genai.GenerativeModel(model or "gemini-1.5-flash",system_instruction=SYSTEM_PROMPT)
    h=[{"role":"user" if x["role"]=="user" else "model","parts":[x["content"]]} for x in hist[-20:]]
    chat=m.start_chat(history=h)
    if img_b64:
        import PIL.Image,io
        img=PIL.Image.open(io.BytesIO(base64.b64decode(img_b64)))
        r=chat.send_message([msg or "Analyze this image.",img])
    else:
        r=chat.send_message(msg,generation_config={"temperature":temp,"max_output_tokens":4096})
    return r.text

def _openai(msg,hist,model,temp,img_b64):
    from openai import OpenAI
    client=OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
    msgs=[{"role":"system","content":SYSTEM_PROMPT}]
    msgs.extend(hist[-20:])
    if img_b64:
        msgs.append({"role":"user","content":[
            {"type":"image_url","image_url":{"url":f"data:image/jpeg;base64,{img_b64}"}},
            {"type":"text","text":msg or "Analyze this image."}]})
        model=model or "gpt-4o"
    else:
        msgs.append({"role":"user","content":msg}); model=model or "gpt-4o-mini"
    r=client.chat.completions.create(model=model,messages=msgs,temperature=temp,max_tokens=4096)
    return r.choices[0].message.content

def search_web(q,n=5):
    try:
        from duckduckgo_search import DDGS
        with DDGS() as d: res=list(d.text(q,max_results=n))
        if not res: return ""
        return "\n\n".join(f"[{i+1}] {r.get('title','')}\n{r.get('body','')}\n{r.get('href','')}" for i,r in enumerate(res))
    except: return ""

def read_file(f):
    name=f.name; ext=Path(name).suffix.lower()
    try:
        if ext in [".txt",".md",".py",".js",".ts",".html",".css",".java",".cpp",".c",
                   ".go",".rs",".php",".sh",".sql",".json",".xml",".yaml",".yml",".csv"]:
            return f.read().decode("utf-8",errors="replace")
        elif ext==".pdf":
            import PyPDF2; r=PyPDF2.PdfReader(f)
            return "\n".join(p.extract_text() or "" for p in r.pages)
        elif ext==".docx":
            from docx import Document; d=Document(f)
            return "\n".join(p.text for p in d.paragraphs if p.text.strip())
        elif ext==".xlsx":
            import openpyxl; wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
            rows=[]
            for ws in wb.worksheets:
                rows.append(f"=== {ws.title} ===")
                for row in ws.iter_rows(max_row=100,values_only=True):
                    rows.append(" | ".join(str(c) if c else "" for c in row))
            return "\n".join(rows)
        else: return f"[File: {name}]"
    except Exception as e: return f"Error: {e}"

FILE_TYPES={"py":"python","js":"javascript","ts":"typescript","jsx":"jsx","tsx":"tsx",
    "html":"html","css":"css","java":"java","cpp":"cpp","c":"c","cs":"csharp",
    "go":"go","rs":"rust","php":"php","rb":"ruby","swift":"swift","kt":"kotlin",
    "sql":"sql","sh":"bash","json":"json","xml":"xml","yaml":"yaml","csv":"csv",
    "md":"markdown","txt":"text","r":"r"}

def gen_file(prompt,ext,provider=None):
    lang=FILE_TYPES.get(ext,ext)
    if ext in ["md","txt","csv","json","xml","yaml"]:
        p=f"Create a complete {lang} file for:\n{prompt}\n\nReturn ONLY raw content."
    else:
        p=f"Write complete production-ready {lang} code for:\n{prompt}\n\nReturn ONLY code in a ```{lang} block."
    raw=ai(p,provider=provider,temperature=0.3)
    for pat in [rf'```{lang}\n(.*?)```',r'```\w*\n(.*?)```',r'```\n(.*?)```']:
        m=re.search(pat,raw,re.DOTALL)
        if m: return m.group(1).strip()
    return raw.strip()

def sidebar():
    t=T()
    with st.sidebar:
        st.markdown(f"""<div style="padding:.8rem 0 1.2rem;text-align:center;">
            <div style="font-size:1.5rem;font-weight:700;color:{t['text']};letter-spacing:-.5px;">CherryAI</div>
            <div style="font-size:.72rem;color:{t['sub']};margin-top:2px;">Personal AI Assistant</div>
        </div>""",unsafe_allow_html=True)

        c1,c2=st.columns(2)
        with c1:
            if st.button("New Chat",use_container_width=True):
                st.session_state.messages=[]; st.session_state.editing_index=None; st.rerun()
        with c2:
            lbl="☀️ Light" if st.session_state.dark_mode else "Dark"
            if st.button(lbl,use_container_width=True):
                st.session_state.dark_mode=not st.session_state.dark_mode; st.rerun()

        st.markdown("---")
        st.markdown(f"<div style='font-size:.82rem;font-weight:600;color:{t['text']};margin-bottom:.4rem;'> Memory Search</div>",unsafe_allow_html=True)
        sq=st.text_input("Search chats",placeholder="Search past messages...",label_visibility="collapsed")
        if sq:
            found=[(i,m) for i,m in enumerate(st.session_state.messages) if sq.lower() in m["content"].lower()]
            if found:
                st.markdown(f"<div style='font-size:.73rem;color:{t['sub']};margin-bottom:.3rem;'>{len(found)} result(s) found</div>",unsafe_allow_html=True)
                for i,m in found[:5]:
                    icon="" if m["role"]=="user" else "🍒"
                    prev=m["content"][:55]+"..." if len(m["content"])>55 else m["content"]
                    st.markdown(f'<div class="srch-result">{icon} {prev}</div>',unsafe_allow_html=True)
            else:
                st.markdown(f"<div style='font-size:.73rem;color:{t['sub']};'>No results found</div>",unsafe_allow_html=True)

        st.markdown("---")
        st.markdown(f"<div style='font-size:.82rem;font-weight:600;color:{t['text']};margin-bottom:.4rem;'> Settings</div>",unsafe_allow_html=True)
        pm={"auto":"Auto-Select","groq":"Groq (FREE)","gemini":"Gemini (FREE)","openai":"OpenAI (Paid)"}
        sel=st.selectbox("Provider",list(pm.keys()),format_func=lambda x:pm[x])
        st.session_state.provider=sel
        if sel=="groq":
            st.session_state.model=st.selectbox("Model",["llama-3.3-70b-versatile","llama-3.1-8b-instant","gemma2-9b-it"])
        elif sel=="gemini":
            st.session_state.model=st.selectbox("Model",["gemini-1.5-flash","gemini-1.5-pro"])
        else:
            st.session_state.model=None
        st.session_state.temperature=st.slider("Creativity",0.0,1.0,st.session_state.temperature,0.1)
        st.session_state.web_search=st.toggle("🌐 Web Search",value=st.session_state.web_search)
        st.markdown("---")
        c1,c2=st.columns(2)
        with c1: st.metric("Messages",len(st.session_state.messages))
        with c2: st.metric("Asked",sum(1 for m in st.session_state.messages if m["role"]=="user"))
        st.markdown("---")
        with st.expander(" API Setup"):
            st.markdown("**Free Groq:**\n1. [console.groq.com](https://console.groq.com)\n2. Create Key\n3. Add to HF secrets:\n```\nGROQ_API_KEY = 'your_key'\n```")

def welcome():
    t=T()
    st.markdown(f"""<div class="welcome-wrap">
        <div class="welcome-t">Welcome to CherryAI 🍒</div>
        <div class="welcome-s">Ask anything · Generate files · Fix code · Analyze images</div>
    </div>""",unsafe_allow_html=True)
    sugg=["Write a Python Flask REST API with CRUD","Explain machine learning simply",
          "Fix my code bugs and optimize it","Create a responsive HTML landing page",
          "Write SQL query for top customers","Who built CherryAI?"]
    cols=st.columns(2)
    for i,s in enumerate(sugg):
        with cols[i%2]:
            if st.button(s[:50]+("..." if len(s)>50 else ""),use_container_width=True,key=f"s{i}"):
                st.session_state["_pre"]=s; st.rerun()

def render_msgs():
    t=T()
    for i,msg in enumerate(st.session_state.messages):
        if msg["role"]=="user":
            if st.session_state.editing_index==i:
                new=st.text_area("Edit",value=msg["content"],key=f"ed{i}",height=80)
                c1,c2,_=st.columns([1,1,4])
                with c1:
                    if st.button("Save",key=f"sv{i}"):
                        st.session_state.messages[i]["content"]=new
                        st.session_state.messages=st.session_state.messages[:i+1]
                        st.session_state.editing_index=None
                        st.session_state["_resend"]=i; st.rerun()
                with c2:
                    if st.button("Cancel",key=f"cn{i}"):
                        st.session_state.editing_index=None; st.rerun()
            else:
                c1,c2=st.columns([0.92,0.08])
                with c1:
                    img_html=""
                    if msg.get("img_b64"):
                        img_html=f'<img src="data:image/jpeg;base64,{msg["img_b64"]}" class="chat-img"/>'
                    st.markdown(f'<div class="user-row"><div class="user-bbl">{img_html}{msg["content"]}</div></div>',unsafe_allow_html=True)
                with c2:
                    st.markdown("<div style='height:8px'></div>",unsafe_allow_html=True)
                    if st.button("✏",key=f"eb{i}",help="Edit this message"):
                        st.session_state.editing_index=i; st.rerun()
        else:
            c1,c2=st.columns([0.045,0.955])
            with c1: st.markdown('<div class="cherry-av">🍒</div>',unsafe_allow_html=True)
            with c2: st.markdown(msg["content"])
    st.markdown('<div id="btm"></div>',unsafe_allow_html=True)
    st.markdown("<script>document.getElementById('btm')?.scrollIntoView({behavior:'smooth'});</script>",unsafe_allow_html=True)

def chat_tab():
    t=T()
    resend=st.session_state.pop("_resend",None)
    if resend is not None:
        msg=st.session_state.messages[resend]
        with st.spinner("Thinking..."):
            hist=[{"role":m["role"],"content":m["content"]} for m in st.session_state.messages[:resend]]
            prov=st.session_state.provider if st.session_state.provider!="auto" else None
            resp=ai(msg["content"],history=hist,provider=prov,model=st.session_state.model,temperature=st.session_state.temperature)
        st.session_state.messages.append({"role":"assistant","content":resp}); st.rerun()

    if st.session_state.messages:
        render_msgs()
        if len(st.session_state.messages)>4:
            st.markdown('<a href="#btm"><button class="scroll-down-btn">↓</button></a>',unsafe_allow_html=True)
    else:
        welcome()

    st.markdown(f"<div style='font-size:.75rem;color:{t['sub']};margin-top:.5rem;margin-bottom:.2rem;'>📎 Attach image (optional)</div>",unsafe_allow_html=True)
    up_img=st.file_uploader("Image",type=["png","jpg","jpeg","gif","webp"],label_visibility="collapsed",key="ci")
    pre=st.session_state.pop("_pre",None)
    user_input=st.chat_input("Message CherryAI...")
    if pre: user_input=pre
    if user_input:
        img_b64=None
        if up_img:
            up_img.seek(0); img_b64=base64.b64encode(up_img.read()).decode("utf-8")
        umsg={"role":"user","content":user_input}
        if img_b64: umsg["img_b64"]=img_b64
        st.session_state.messages.append(umsg)
        with st.spinner("Thinking..."):
            sc=""
            if st.session_state.web_search:
                sc=search_web(user_input)
                if sc: sc=f"\n\n[WEB SEARCH]\n{sc}\n[END]\n\n"
            fm=sc+user_input
            hist=[{"role":m["role"],"content":m["content"]} for m in st.session_state.messages[:-1]]
            prov=st.session_state.provider if st.session_state.provider!="auto" else None
            resp=ai(fm,history=hist,provider=prov,model=st.session_state.model,temperature=st.session_state.temperature,img_b64=img_b64)
        if sc: resp="🌐 *Web search used*\n\n"+resp
        st.session_state.messages.append({"role":"assistant","content":resp}); st.rerun()

def generate_tab():
    t=T()
    st.markdown(f'<div class="section-card"><div class="sec-title">⚡ Generate Any File</div><div class="sec-sub">Describe what you want — CherryAI writes the complete file instantly</div></div>',unsafe_allow_html=True)
    c1,c2=st.columns([3,1])
    with c1:
        prompt=st.text_area("What should the file do?",height=120,placeholder="e.g. Python FastAPI server with JWT authentication and user CRUD endpoints")
    with c2:
        st.markdown("<div style='height:8px'></div>",unsafe_allow_html=True)
        ext=st.selectbox("Type",list(FILE_TYPES.keys()))
        fname=st.text_input("Filename",placeholder=f"output.{ext}")
    if st.button("Generate File",use_container_width=True,type="primary"):
        if not prompt.strip(): st.warning("Please describe what you want.")
        else:
            with st.spinner(f"Generating {ext.upper()}..."):
                prov=st.session_state.provider if st.session_state.provider!="auto" else None
                content=gen_file(prompt,ext,provider=prov)
                fn=fname if fname else f"cherryai.{ext}"
                if not fn.endswith(f".{ext}"): fn+=f".{ext}"
            lines=content.count('\n')+1
            st.markdown(f'<div style="background:{t["card"]};border:1px solid {t["cborder"]};border-radius:10px;padding:.7rem 1rem;margin:.5rem 0;font-size:.85rem;color:{t["text"]};"><b>✅ {fn}</b> &nbsp;·&nbsp; <span style="color:{t["sub"]};">{lines} lines · {len(content)} chars</span></div>',unsafe_allow_html=True)
            st.code(content,language=FILE_TYPES.get(ext,ext))
            st.download_button(f" Download {fn}",data=content,file_name=fn,mime="text/plain",use_container_width=True)

def code_tab():
    t=T()
    st.markdown(f'<div class="section-card"><div class="sec-title"> Code Tools</div><div class="sec-sub">Fix bugs · Explain · Optimize · Add documentation to any code</div></div>',unsafe_allow_html=True)
    c1,c2=st.columns([4,1])
    with c1:
        code=st.text_area("Paste your code",height=280,placeholder="Paste any code here...")
    with c2:
        langs=["python","javascript","typescript","java","cpp","c","go","rust","php","ruby","sql","bash","html","css"]
        lang=st.selectbox("Language",langs)
        task=st.selectbox("Task",["fix","explain","optimize","document"])
        info={"fix":("","Fix Bugs"),"explain":("","Explain"),"optimize":("⚡","Optimize"),"document":("","Document")}
        ic,lb=info[task]
        st.markdown(f'<div style="background:{t["card"]};border:1px solid {t["cborder"]};border-radius:8px;padding:.5rem;text-align:center;margin-top:.3rem;"><div style="font-size:1.1rem;">{ic}</div><div style="font-size:.73rem;color:{t["sub"]};">{lb}</div></div>',unsafe_allow_html=True)
    if st.button("Run",use_container_width=True,type="primary"):
        if not code.strip(): st.warning("Paste some code first.")
        else:
            tp={"fix":f"Fix ALL bugs in this {lang} code. Return corrected code + explanation:\n\n```{lang}\n{code}\n```",
                "explain":f"Explain this {lang} code step by step:\n\n```{lang}\n{code}\n```",
                "optimize":f"Optimize this {lang} code. Return improved code + changes:\n\n```{lang}\n{code}\n```",
                "document":f"Add comprehensive comments/docstrings:\n\n```{lang}\n{code}\n```"}
            with st.spinner("Analyzing..."):
                prov=st.session_state.provider if st.session_state.provider!="auto" else None
                res=ai(tp[task],provider=prov)
            st.markdown(f'<div style="font-size:.85rem;font-weight:600;color:{t["text"]};margin:1rem 0 .5rem;padding-bottom:.3rem;border-bottom:1px solid {t["hr"]};">Result</div>',unsafe_allow_html=True)
            st.markdown(res)

def analyze_tab():
    t=T()
    st.markdown(f'<div class="section-card"><div class="sec-title">📎 Analyze Any File</div><div class="sec-sub">Upload PDF · Word · Excel · Code · CSV · Images — ask anything</div></div>',unsafe_allow_html=True)
    c1,c2=st.columns(2)
    with c1:
        udoc=st.file_uploader("Upload document",type=["pdf","docx","txt","py","js","html","css","csv","xlsx","json","xml","yaml","md","java","cpp","sql"])
    with c2:
        uimg=st.file_uploader("Upload image to analyze",type=["png","jpg","jpeg","webp","gif"])
    q=st.text_input("Your question",value="Summarize and analyze this in detail.")
    if st.button("Analyze",use_container_width=True,type="primary"):
        prov=st.session_state.provider if st.session_state.provider!="auto" else None
        if uimg:
            with st.spinner(f"Analyzing image..."):
                uimg.seek(0); ib64=base64.b64encode(uimg.read()).decode("utf-8")
                res=ai(q,provider=prov,img_b64=ib64)
            st.success(f"✅ Analyzed: {uimg.name}")
            uimg.seek(0); st.image(uimg,width=300)
            st.markdown("### Result"); st.markdown(res)
        elif udoc:
            with st.spinner(f"Reading {udoc.name}..."):
                content=read_file(udoc)
                res=ai(f"File: {udoc.name}\n\nContent:\n{content[:8000]}\n\nQuestion: {q}",provider=prov)
            st.success(f"✅ Analyzed: {udoc.name}")
            st.markdown("### Result"); st.markdown(res)
        else:
            st.warning("Upload a file or image first.")

def footer():
    st.markdown('<div class="cherry-footer">🍒 &nbsp; © 2026 <span>CherryAI</span> &nbsp;·&nbsp; Built by Sai Vignesh &nbsp;·&nbsp; All rights reserved</div>',unsafe_allow_html=True)

def main():
    t=T()
    sidebar()
    st.markdown(f'<div class="cherry-hdr"><span class="cherry-hdr-t">CherryAI</span><span class="cherry-hdr-s">Personal AI Assistant</span></div>',unsafe_allow_html=True)
    tabs=st.tabs([" Chat"," Generate File"," Code Tools"," Analyze File"])
    with tabs[0]: chat_tab()
    with tabs[1]: generate_tab()
    with tabs[2]: code_tab()
    with tabs[3]: analyze_tab()
    footer()

if __name__=="__main__":
    main()
